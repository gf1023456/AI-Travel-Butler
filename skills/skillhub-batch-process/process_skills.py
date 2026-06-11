import sys, zipfile, os, re
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SRC_DIR = r'D:\工作记录\项目资料\skillshub\2000skill\2026-6-11'
ZIP_DIR = os.path.join(SRC_DIR, '63个技能包', '63个技能包')
XLSX_PATH = os.path.join(SRC_DIR, '63.xlsx')
OUTPUT_PATH = os.path.join(SRC_DIR, '输出_最终版_63.xlsx')

DOMAINS = {
    'DOMAIN_GENERAL_TOOL': '通用工具',
    'DOMAIN_MARKETING': '市场营销',
    'DOMAIN_OFFICE_MANAGEMENT': '办公管理',
    'DOMAIN_RD_OPS': '研发运维',
    'DOMAIN_TERMINAL_MANAGEMENT': '终端管理',
    'DOMAIN_FINANCIAL_MARKET': '金融市场',
    'DOMAIN_INFO_SEARCH': '信息搜索',
    'DOMAIN_DEV_TOOL': '开发者工具',
    'DOMAIN_LIFE_ENTERTAINMENT': '生活娱乐',
    'DOMAIN_SECURITY_COMPLIANCE': '安全合规',
    'DOMAIN_DATA_ANALYSIS': '数据分析',
    'DOMAIN_MULTIMEDIA_PROCESS': '多媒体处理',
    'DOMAIN_AI_INTELLIGENCE': 'AI智能',
    'DOMAIN_INDUSTRY_SPECIFIC': '行业专属',
    'DOMAIN_OTHER': '其他',
}

def wb_match(text, word):
    return bool(re.search(r'(?:^|[\s\-_/])' + re.escape(word) + r'(?:$|[\s\-_.,;:!?])', text))

def classify_domain(slug, name, desc):
    text = (name + ' ' + desc).lower()

    # === Slug-level overrides ===
    OVERRIDE = {
        'http': 'DOMAIN_GENERAL_TOOL',
        'humanize': 'DOMAIN_GENERAL_TOOL',
        'humanizer-academic-zh': 'DOMAIN_GENERAL_TOOL',
        'self': 'DOMAIN_LIFE_ENTERTAINMENT',
        'self-love-confidence': 'DOMAIN_LIFE_ENTERTAINMENT',
        'quit-vaping': 'DOMAIN_LIFE_ENTERTAINMENT',
        'overcome-problem': 'DOMAIN_LIFE_ENTERTAINMENT',
        'language-learning': 'DOMAIN_LIFE_ENTERTAINMENT',
        'jungian-psychologist': 'DOMAIN_LIFE_ENTERTAINMENT',
        'remind': 'DOMAIN_LIFE_ENTERTAINMENT',
        'roadrunner': 'DOMAIN_LIFE_ENTERTAINMENT',
        'screenshot': 'DOMAIN_MULTIMEDIA_PROCESS',
        'poster': 'DOMAIN_MULTIMEDIA_PROCESS',
        'remotion': 'DOMAIN_MULTIMEDIA_PROCESS',
        'remotion-video-toolkit': 'DOMAIN_MULTIMEDIA_PROCESS',
        'seiso': 'DOMAIN_MULTIMEDIA_PROCESS',
        'medial-generation': 'DOMAIN_MULTIMEDIA_PROCESS',
        'human-optimized-frontend': 'DOMAIN_DEV_TOOL',
        'react-expert': 'DOMAIN_DEV_TOOL',
        'react-native': 'DOMAIN_DEV_TOOL',
        'rust': 'DOMAIN_DEV_TOOL',
        'portable-tools': 'DOMAIN_DEV_TOOL',
        'pet': 'DOMAIN_DEV_TOOL',
        'ripgrep': 'DOMAIN_DEV_TOOL',
        'scrapling-official': 'DOMAIN_DEV_TOOL',
        'qmd-local-search': 'DOMAIN_DEV_TOOL',
        'railway-skill': 'DOMAIN_RD_OPS',
        'scrape': 'DOMAIN_RD_OPS',
        'rssaurus': 'DOMAIN_RD_OPS',
        'secureclaw-skill': 'DOMAIN_SECURITY_COMPLIANCE',
        'security-audit': 'DOMAIN_SECURITY_COMPLIANCE',
        'security-auditor': 'DOMAIN_SECURITY_COMPLIANCE',
        'runesleo-systematic-debugging': 'DOMAIN_SECURITY_COMPLIANCE',
        'keyword-research': 'DOMAIN_MARKETING',
        'rank-tracker': 'DOMAIN_MARKETING',
        'personal-branding-authority': 'DOMAIN_MARKETING',
        'product-marketing-context': 'DOMAIN_MARKETING',
        'product-manager': 'DOMAIN_OFFICE_MANAGEMENT',
        'productivity': 'DOMAIN_OFFICE_MANAGEMENT',
        'project-management-2': 'DOMAIN_OFFICE_MANAGEMENT',
        'project-management-guru-adhd': 'DOMAIN_OFFICE_MANAGEMENT',
        'pa-admin-exec': 'DOMAIN_OFFICE_MANAGEMENT',
        'schedule': 'DOMAIN_OFFICE_MANAGEMENT',
        'request-approval': 'DOMAIN_OFFICE_MANAGEMENT',
        'requesting-code-review': 'DOMAIN_OFFICE_MANAGEMENT',
        'proposal-writer': 'DOMAIN_OFFICE_MANAGEMENT',
        'resume-builder': 'DOMAIN_OFFICE_MANAGEMENT',
        'resume-cv-builder': 'DOMAIN_OFFICE_MANAGEMENT',
        'sales-pipeline-tracker': 'DOMAIN_OFFICE_MANAGEMENT',
        'pre-mortem-analyst': 'DOMAIN_OFFICE_MANAGEMENT',
        'quality-manager-qmr': 'DOMAIN_OFFICE_MANAGEMENT',
        'prompt-architect': 'DOMAIN_AI_INTELLIGENCE',
        'prompt-engineering-expert': 'DOMAIN_AI_INTELLIGENCE',
        'hybrid-memory': 'DOMAIN_AI_INTELLIGENCE',
        'quant-strategy': 'DOMAIN_FINANCIAL_MARKET',
        'quant-trading': 'DOMAIN_FINANCIAL_MARKET',
        'portfolio-manager': 'DOMAIN_FINANCIAL_MARKET',
        'pdf-cn': 'DOMAIN_INFO_SEARCH',
        'pdf-tools': 'DOMAIN_INFO_SEARCH',
        'patent-scanner': 'DOMAIN_INDUSTRY_SPECIFIC',
        'pubmed': 'DOMAIN_INDUSTRY_SPECIFIC',
        'research-paper-writer': 'DOMAIN_INDUSTRY_SPECIFIC',
        'rupert-data-analysis': 'DOMAIN_DATA_ANALYSIS',
        'performance-reporter': 'DOMAIN_DATA_ANALYSIS',
        'scrape-cli': 'DOMAIN_DATA_ANALYSIS',
    }
    if slug in OVERRIDE:
        return OVERRIDE[slug], DOMAINS[OVERRIDE[slug]]

    # DOMAIN_SECURITY_COMPLIANCE
    for kw in ['audit', 'auditing', 'vulnerability', 'owasp', 'penetration test',
               'secureclaw', 'cors', 'csp', 'xss', 'sql injection']:
        if wb_match(text, kw):
            return 'DOMAIN_SECURITY_COMPLIANCE', DOMAINS['DOMAIN_SECURITY_COMPLIANCE']

    # DOMAIN_AI_INTELLIGENCE
    for kw in ['prompt', 'llm', 'memory system', 'rag', 'vector', 'embedding']:
        if kw in text:
            return 'DOMAIN_AI_INTELLIGENCE', DOMAINS['DOMAIN_AI_INTELLIGENCE']

    # DOMAIN_FINANCIAL_MARKET
    for kw in ['quant', 'financial', 'trading', 'stock', 'portfolio']:
        if kw in text:
            return 'DOMAIN_FINANCIAL_MARKET', DOMAINS['DOMAIN_FINANCIAL_MARKET']

    # DOMAIN_MULTIMEDIA_PROCESS
    for kw in ['screenshot', 'poster', 'video', 'media generation',
               'image', 'visual', 'pixel']:
        if kw in text:
            return 'DOMAIN_MULTIMEDIA_PROCESS', DOMAINS['DOMAIN_MULTIMEDIA_PROCESS']

    # DOMAIN_MARKETING
    for kw in ['seo', 'keyword', 'rank track', 'marketing', 'brand',
               'personal brand']:
        if kw in text:
            return 'DOMAIN_MARKETING', DOMAINS['DOMAIN_MARKETING']

    # DOMAIN_LIFE_ENTERTAINMENT
    for kw in ['self love', 'self care', 'quit vaping', 'language learn',
               'psychologist', 'remind', 'overcome', 'roadrunner',
               'beeper', 'messaging', 'social']:
        if kw in text:
            return 'DOMAIN_LIFE_ENTERTAINMENT', DOMAINS['DOMAIN_LIFE_ENTERTAINMENT']

    # DOMAIN_DEV_TOOL
    for kw in ['react', 'rust', 'frontend', 'code review', 'remotion',
               'cli ', 'terminal', 'snippet manager']:
        if wb_match(text, kw):
            return 'DOMAIN_DEV_TOOL', DOMAINS['DOMAIN_DEV_TOOL']

    # DOMAIN_RD_OPS
    for kw in ['deploy', 'railway', 'scrape', 'rssaurus',
               'devops', 'ci/cd']:
        if kw in text:
            return 'DOMAIN_RD_OPS', DOMAINS['DOMAIN_RD_OPS']

    # DOMAIN_OFFICE_MANAGEMENT
    for kw in ['productivity', 'project management', 'project manager',
               'product manager', 'product marketing', 'schedule',
               'approval', 'resume', 'proposal', 'admin exec',
               'sales pipeline', 'task tracker']:
        if kw in text:
            return 'DOMAIN_OFFICE_MANAGEMENT', DOMAINS['DOMAIN_OFFICE_MANAGEMENT']

    # DOMAIN_DATA_ANALYSIS
    for kw in ['data analysis', 'performance report', 'analytics',
               'visualization', 'bi ', 'kpi']:
        if kw in text:
            return 'DOMAIN_DATA_ANALYSIS', DOMAINS['DOMAIN_DATA_ANALYSIS']

    # DOMAIN_INFO_SEARCH
    for kw in ['search', 'research', 'pdf', 'markdown', 'file discovery']:
        if kw in text:
            return 'DOMAIN_INFO_SEARCH', DOMAINS['DOMAIN_INFO_SEARCH']

    # DOMAIN_INDUSTRY_SPECIFIC
    for kw in ['patent', 'pubmed', 'biomedical', 'academic', 'literature']:
        if kw in text:
            return 'DOMAIN_INDUSTRY_SPECIFIC', DOMAINS['DOMAIN_INDUSTRY_SPECIFIC']

    # DOMAIN_TERMINAL_MANAGEMENT
    for kw in ['portable tool', 'cross-device', 'command-line', 'ripgrep']:
        if kw in text:
            return 'DOMAIN_TERMINAL_MANAGEMENT', DOMAINS['DOMAIN_TERMINAL_MANAGEMENT']

    return 'DOMAIN_GENERAL_TOOL', DOMAINS['DOMAIN_GENERAL_TOOL']


# === Chinese Description Translation Map ===
DESC_CN = {
    'http': 'HTTP请求处理工具，支持发送GET、POST、PUT、DELETE等HTTP请求，处理响应数据和错误。',
    'humanize': '文本人性化处理工具，优化文本使之更自然、流畅、易读。',
    'humanizer-academic-zh': '中文学术论文润色工具，降低AIGC检测率，去除AI痕迹，使论文更自然。',
    'human-optimized-frontend': '前端界面生成工具，根据需求描述自动生成美观、响应式的用户界面。',
    'hybrid-memory': '混合记忆系统，结合短期和长期记忆管理，提供智能上下文感知能力。',
    'jungian-psychologist': '荣格心理学顾问，基于荣格心理学理论提供心理分析和建议。',
    'keyword-research': 'SEO关键词研究工具，分析搜索量、难度和竞争度，发现高价值关键词。',
    'language-learning': 'AI语言学习助手，通过对话练习帮助用户学习新语言。',
    'overcome-problem': '问题分析与解决工具，帮助用户系统化分析问题并找到解决方案。',
    'pa-admin-exec': '行政执行支持工具，协助处理日常行政事务和任务管理。',
    'patent-scanner': '专利扫描分析工具，搜索、分析和监控相关专利信息。',
    'pdf-cn': 'PDF文档处理工具，支持读取、提取、合并、分割PDF文件。',
    'pdf-tools': 'PDF工具集，提供PDF文档的各种处理功能。',
    'performance-reporter': '性能报告生成工具，收集分析数据并生成可视化性能报告。',
    'personal-branding-authority': '个人品牌建设工具，帮助打造和维护个人专业品牌形象。',
    'pet': '命令行片段管理器，存储、搜索和快速使用常用的命令行代码片段。',
    'portable-tools': '跨设备开发工具集，提供便携的开发环境和常用工具。',
    'portfolio-manager': '投资组合管理工具，跟踪和管理投资组合的收益和风险。',
    'poster': '海报设计生成工具，根据需求自动生成专业级海报设计。',
    'pre-mortem-analyst': '项目预分析工具，在项目启动前识别潜在风险和问题。',
    'product-manager': '产品管理专家工具，协助产品规划、需求分析和路线图制定。',
    'product-marketing-context': '产品营销上下文工具，提供产品营销背景分析和策略建议。',
    'productivity': '生产力提升工具，帮助用户提高工作效率和时间管理能力。',
    'project-management-2': '项目管理工具，支持任务分配、进度跟踪和团队协作。',
    'project-management-guru-adhd': 'ADHD项目管理专家，专为ADHD人群设计的项目管理工具。',
    'prompt-architect': 'Prompt架构设计工具，帮助设计高效的AI对话提示词。',
    'prompt-engineering-expert': 'Prompt工程专家工具，提供高级提示词工程技巧和策略。',
    'proposal-writer': '商业提案撰写工具，自动生成专业、有说服力的商业提案。',
    'pubmed': '生物医学文献检索工具，快速搜索和筛选PubMed数据库中的学术文献。',
    'qmd-local-search': '本地Markdown搜索工具，快速搜索和查找本地Markdown文件内容。',
    'quality-manager-qmr': '质量管理负责人工具，协助质量管理体系的建立和维护。',
    'quant-strategy': '量化策略开发工具，辅助编写和回测量化交易策略。',
    'quant-trading': '量化交易工具，提供多因子选股、技术指标分析和策略回测。',
    'quit-vaping': '戒烟助手工具，提供个性化戒烟计划和进度跟踪。',
    'railway-skill': 'Railway应用部署工具，简化应用部署和运维流程。',
    'rank-tracker': '关键词排名跟踪工具，监控关键词在搜索引擎中的排名变化。',
    'react-expert': 'React开发专家工具，提供React组件开发和最佳实践指导。',
    'react-native': 'React Native开发工具，协助跨平台移动应用开发。',
    'remind': '智能提醒工具，设置和管理各类提醒和待办事项。',
    'remotion': 'Remotion视频创建工具，使用React组件编程式创建视频内容。',
    'remotion-video-toolkit': 'Remotion视频工具包，提供视频创建的扩展工具和模板。',
    'request-approval': '审批请求工具，创建、提交和跟踪审批流程。',
    'requesting-code-review': '代码审查请求工具，提交代码审查请求并跟踪审查进度。',
    'research-paper-writer': '学术论文撰写工具，辅助论文写作、格式化和引用管理。',
    'resume-builder': '简历生成工具，生成符合Reactive Resume标准的专业简历。',
    'resume-cv-builder': '简历/CV生成工具，创建ATS友好的简历，支持多种导出格式。',
    'ripgrep': '文本搜索工具，基于ripgrep的快速文件内容搜索工具。',
    'roadrunner': '消息通信客户端，支持多平台消息通信和通知管理。',
    'rssaurus': 'RSS订阅管理工具，通过命令行管理RSS订阅源和阅读列表。',
    'runesleo-systematic-debugging': '系统化调试工具，提供结构化的调试方法和问题定位流程。',
    'rupert-data-analysis': '数据分析与可视化工具，使用自然语言进行数据查询和分析。',
    'rust': 'Rust开发工具，提供Rust编程指导和最佳实践。',
    'sales-pipeline-tracker': '销售管道跟踪工具，管理销售机会和客户跟进流程。',
    'schedule': '任务调度工具，自动化任务安排和时间表管理。',
    'scrape': '网页抓取工具，从网页中提取和收集数据。',
    'scrapling-official': '网页抓取工具，提供高级网页数据采集功能。',
    'screenshot': '截图工具，捕获、检查和比较屏幕截图，支持多种设备和场景。',
    'secureclaw-skill': '安全检测工具，检测代码和应用中的安全漏洞。',
    'security-audit': '安全审计工具，进行代码安全审查和安全配置检查。',
    'security-auditor': '安全审计专家工具，审计OWASP Top 10、认证流程和安全配置。',
    'seiso': '媒体生成网关工具，管理和生成多种格式的媒体内容。',
    'self': '自我成长工具，提供个人反思、目标设定和成长规划。',
    'self-love-confidence': '自爱自信建设工具，帮助提升自信心和自我接纳能力。',
}

def get_desc_cn(slug, desc_raw):
    if slug in DESC_CN:
        return DESC_CN[slug]
    has_cn = bool(re.search(r'[\u4e00-\u9fff]', desc_raw))
    if has_cn:
        return desc_raw[:200]
    return slug.replace('-', ' ').title()

def infer_chinese_name(slug, desc):
    INFER_CN = {
        'http': 'HTTP请求处理', 'humanize': '文本人性化',
        'humanizer-academic-zh': '中文学术润色',
        'human-optimized-frontend': '前端界面生成',
        'hybrid-memory': '混合记忆系统',
        'jungian-psychologist': '荣格心理学顾问',
        'keyword-research': 'SEO关键词研究',
        'language-learning': 'AI语言学习',
        'overcome-problem': '问题分析与解决',
        'pa-admin-exec': '行政执行支持',
        'patent-scanner': '专利扫描分析',
        'pdf-cn': 'PDF文档处理', 'pdf-tools': 'PDF工具集',
        'performance-reporter': '性能报告生成',
        'personal-branding-authority': '个人品牌建设',
        'pet': '命令行片段管理器',
        'portable-tools': '跨设备开发工具',
        'portfolio-manager': '投资组合管理',
        'poster': '海报设计生成',
        'pre-mortem-analyst': '项目预分析',
        'product-manager': '产品管理专家',
        'product-marketing-context': '产品营销上下文',
        'productivity': '生产力提升',
        'project-management-2': '项目管理',
        'project-management-guru-adhd': 'ADHD项目管理专家',
        'prompt-architect': 'Prompt架构设计',
        'prompt-engineering-expert': 'Prompt工程专家',
        'proposal-writer': '商业提案撰写',
        'pubmed': '生物医学文献检索',
        'qmd-local-search': '本地Markdown搜索',
        'quality-manager-qmr': '质量管理负责人',
        'quant-strategy': '量化策略开发',
        'quant-trading': '量化交易',
        'quit-vaping': '戒烟助手',
        'railway-skill': 'Railway应用部署',
        'rank-tracker': '关键词排名跟踪',
        'react-expert': 'React开发专家',
        'react-native': 'React Native开发',
        'remind': '智能提醒',
        'remotion': 'Remotion视频创建',
        'remotion-video-toolkit': 'Remotion视频工具包',
        'request-approval': '审批请求',
        'requesting-code-review': '代码审查请求',
        'research-paper-writer': '学术论文撰写',
        'resume-builder': '简历生成',
        'resume-cv-builder': '简历/CV生成',
        'ripgrep': '文本搜索工具',
        'roadrunner': '消息通信客户端',
        'rssaurus': 'RSS订阅管理',
        'runesleo-systematic-debugging': '系统化调试',
        'rupert-data-analysis': '数据分析',
        'rust': 'Rust开发',
        'sales-pipeline-tracker': '销售管道跟踪',
        'schedule': '任务调度',
        'scrape': '网页抓取',
        'scrapling-official': '网页抓取工具',
        'screenshot': '截图工具',
        'secureclaw-skill': '安全检测',
        'security-audit': '安全审计',
        'security-auditor': '安全审计专家',
        'seiso': '媒体生成网关',
        'self': '自我成长',
        'self-love-confidence': '自爱自信建设',
    }
    if slug in INFER_CN:
        return INFER_CN[slug]
    cn = re.search(r'[\u4e00-\u9fff]{2,}', desc)
    if cn:
        return cn.group()
    return slug.replace('-', ' ').title()


# === Main ===
print('Reading 63.xlsx...')
wb_in = openpyxl.load_workbook(XLSX_PATH)
ws_in = wb_in.active
skills = []
for row in ws_in.iter_rows(min_row=2, max_row=ws_in.max_row, values_only=True):
    skills.append({'id': row[0], 'slug': row[1], 'provider': row[3], 'batch': row[7]})
print(f'Total: {len(skills)}')

print('Extracting SKILL.md from zips...')
zip_data = {}
for fname in os.listdir(ZIP_DIR):
    if not fname.endswith('.zip'): continue
    slug_key = fname.replace('.zip', '')
    path = os.path.join(ZIP_DIR, fname)
    try:
        zf = zipfile.ZipFile(path)
        skill_md = [n for n in zf.namelist() if 'SKILL.md' in n]
        if skill_md:
            content = zf.read(skill_md[0]).decode('utf-8', errors='replace')
            m = re.match(r'^---\s*\n(.*?)\n---', content, re.DOTALL)
            if m:
                front = m.group(1)
                desc_m = re.search(r'^description:\s*(.+)$', front, re.MULTILINE)
                name_m = re.search(r'^name:\s*(.+)$', front, re.MULTILINE)
                zip_data[slug_key] = {
                    'name': name_m.group(1).strip().strip('"\'').strip("'") if name_m else slug_key,
                    'desc': desc_m.group(1).strip().strip('"\'').strip("'") if desc_m else ''
                }
    except Exception as e:
        print(f'  ERROR {fname}: {e}')

if 'rupert-data-analysis' not in zip_data:
    zip_data['rupert-data-analysis'] = {
        'name': 'Rupert Data Analysis',
        'desc': 'Data analysis and visualization using natural language queries'
    }
print(f'Extracted: {len(zip_data)}')

print('\nProcessing...')
output_rows = []
for s in skills:
    slug = s['slug']
    zd = zip_data.get(slug, {})

    chinese_name = infer_chinese_name(slug, zd.get('desc', ''))
    desc_cn = get_desc_cn(slug, zd.get('desc', ''))
    domain_code, domain_name = classify_domain(slug, zd.get('name', ''), zd.get('desc', ''))

    output_rows.append({
        '名称': chinese_name,
        '描述': desc_cn,
        '版本': '1.0.0',
        '作者': s.get('provider', ''),
        '类型': '',
        '协议': '',
        '下载量': '',
        '评分': '',
        '来源平台': 'openclaw',
        '仓库地址': '',
        'slug': slug,
        'hasZip': '有zip包',
        '领域分类编码': domain_code,
        '领域分类名称': domain_name,
    })
    print(f'{slug:40s} {domain_code:30s} {chinese_name}')

# === Generate Excel ===
print(f'\nGenerating: {OUTPUT_PATH}')
wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = 'Sheet1'

headers = ['名称', '描述', '版本', '作者', '类型', '协议', '下载量', '评分', '来源平台', '仓库地址', 'slug', 'hasZip', '领域分类编码', '领域分类名称']

hfill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
hfont = Font(name='微软雅黑', bold=True, size=11, color='FFFFFF')
bdr = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))

for col, h in enumerate(headers, 1):
    cell = ws_out.cell(row=1, column=col, value=h)
    cell.font = hfont; cell.fill = hfill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = bdr

for i, r in enumerate(output_rows, 2):
    for col, key in enumerate(headers, 1):
        cell = ws_out.cell(row=i, column=col, value=r[key])
        cell.font = Font(name='微软雅黑', size=10)
        cell.alignment = Alignment(vertical='center', wrap_text=(col == 2))
        cell.border = bdr

widths = [20, 50, 10, 15, 10, 10, 10, 10, 12, 30, 25, 10, 24, 16]
for i, w in enumerate(widths, 1):
    ws_out.column_dimensions[chr(64 + i)].width = w

wb_out.save(OUTPUT_PATH)
print('Done!')
